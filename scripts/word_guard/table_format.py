# -*- coding: utf-8 -*-
"""Academic three-line table formatter for scholarly DOCX documents.

Default presentation contract:
1. Classify tables before formatting. Layout/figure-container tables are never
   converted automatically.
2. Clear the table's previous visual style first so style-based borders/shading
   cannot override the requested three-line appearance.
3. Apply exactly three horizontal rules: table top, header bottom, table bottom.
4. Center cell content horizontally and vertically.
5. Prefer a baseline manuscript's border weights when available.
6. Do not silently change fonts, bolding, widths, merges, or cell text.
"""
from __future__ import annotations

from dataclasses import dataclass, asdict
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


@dataclass
class ThreeLineSpec:
    top_sz: str = '12'       # eighths of a point -> 1.5 pt
    header_sz: str = '6'     # 0.75 pt
    bottom_sz: str = '12'    # 1.5 pt
    color: str = '000000'
    source: str = 'plugin_default'

    def to_dict(self):
        return asdict(self)


def _has_drawing(table) -> bool:
    return table._tbl.find('.//' + qn('w:drawing')) is not None or table._tbl.find('.//' + qn('w:pict')) is not None


def _table_text(table) -> str:
    return '\n'.join(cell.text for row in table.rows for cell in row.cells).strip()


def classify_table(table, preceding_text: str = '') -> dict:
    """Classify a Word table conservatively as data/layout/unknown."""
    text = _table_text(table)
    rows, cols = len(table.rows), len(table.columns)
    prefix = (preceding_text or '').strip().lower()

    if _has_drawing(table):
        return {'kind': 'layout', 'confidence': 'high', 'reason': 'contains drawing/picture'}
    if prefix.startswith('table ') or prefix.startswith('table\t') or prefix.startswith('表'):
        return {'kind': 'data', 'confidence': 'high', 'reason': 'preceded by table caption'}
    # Typical layout containers are tiny and mostly empty.
    nonempty = sum(1 for row in table.rows for cell in row.cells if cell.text.strip())
    total = max(1, rows * cols)
    if rows <= 2 and cols <= 3 and (not text or nonempty / total < 0.5):
        return {'kind': 'layout', 'confidence': 'medium', 'reason': 'small sparse layout-like table'}
    # Conservative data heuristic: at least 2x2 and substantive textual/numeric content.
    if rows >= 2 and cols >= 2 and len(text) >= 20:
        return {'kind': 'data', 'confidence': 'medium', 'reason': 'multi-row/column content table'}
    return {'kind': 'unknown', 'confidence': 'low', 'reason': 'insufficient evidence for safe auto-formatting'}


def _preceding_paragraph_text(table):
    el = table._tbl.getprevious()
    while el is not None:
        if el.tag == qn('w:p'):
            texts = el.findall('.//' + qn('w:t'))
            return ''.join(t.text or '' for t in texts)
        el = el.getprevious()
    return ''


def _edge_sz(container, edge):
    if container is None:
        return None
    el = container.find(qn(f'w:{edge}'))
    if el is None or el.get(qn('w:val')) in (None, 'none', 'nil'):
        return None
    return el.get(qn('w:sz'))


def infer_spec_from_table(table) -> ThreeLineSpec | None:
    """Infer a three-line border specification from an existing baseline table."""
    if not table.rows:
        return None
    tblBorders = table._tbl.tblPr.find(qn('w:tblBorders'))
    top = _edge_sz(tblBorders, 'top')
    bottom = _edge_sz(tblBorders, 'bottom')

    # Prefer row-cell borders; this also recognizes the safer split-table form.
    first = table.rows[0].cells[0]._tc.get_or_add_tcPr().find(qn('w:tcBorders'))
    last = table.rows[-1].cells[0]._tc.get_or_add_tcPr().find(qn('w:tcBorders'))
    top = _edge_sz(first, 'top') or top
    header = _edge_sz(first, 'bottom')
    bottom = _edge_sz(last, 'bottom') or bottom

    if top and header and bottom:
        return ThreeLineSpec(top, header, bottom, '000000', 'baseline')
    return None


def infer_spec_from_baseline(baseline_path: str, table_index: int | None = None) -> ThreeLineSpec | None:
    doc = Document(baseline_path)
    if table_index is not None and 0 <= table_index < len(doc.tables):
        spec = infer_spec_from_table(doc.tables[table_index])
        if spec:
            return spec
    # Fall back to the first confidently data-like table with a recognizable spec.
    for t in doc.tables:
        c = classify_table(t, _preceding_paragraph_text(t))
        if c['kind'] == 'data':
            spec = infer_spec_from_table(t)
            if spec:
                return spec
    return None


def _set_cell_border(cell, edge: str, sz: str, color='000000', val='single'):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn('w:tcBorders'))
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    old = tcBorders.find(qn(f'w:{edge}'))
    if old is not None:
        tcBorders.remove(old)
    el = OxmlElement(f'w:{edge}')
    el.set(qn('w:val'), val)
    el.set(qn('w:sz'), sz)
    el.set(qn('w:space'), '0')
    el.set(qn('w:color'), color)
    tcBorders.append(el)


def _remove_all(parent, tag: str):
    for old in list(parent.findall(qn(tag))):
        parent.remove(old)


def clear_original_table_style(table):
    """Clear visual table styling before applying the academic table format.

    This intentionally removes style-driven and direct borders/shading while
    preserving widths, merges, cell margins, row heights, text, and run fonts.
    """
    tblPr = table._tbl.tblPr
    for tag in ('w:tblStyle', 'w:tblLook', 'w:tblBorders', 'w:shd'):
        _remove_all(tblPr, tag)

    seen_cells = set()
    for row in table.rows:
        for cell in row.cells:
            # Merged cells can appear more than once through python-docx.
            tc_key = cell._tc
            if tc_key in seen_cells:
                continue
            seen_cells.add(tc_key)
            tcPr = cell._tc.get_or_add_tcPr()
            _remove_all(tcPr, 'w:tcBorders')
            _remove_all(tcPr, 'w:shd')


def _center_cell_content(table):
    """Center table cell content horizontally and vertically."""
    seen_cells = set()
    for row in table.rows:
        for cell in row.cells:
            tc_key = cell._tc
            if tc_key in seen_cells:
                continue
            seen_cells.add(tc_key)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def make_three_line_table(table, spec: ThreeLineSpec | None = None):
    """Clear the existing table style, then apply the default three-line table."""
    spec = spec or ThreeLineSpec()
    if not table.rows:
        return {
            'rows': 0,
            'cols': 0,
            'spec': spec.to_dict(),
            'original_style_cleared': False,
            'content_alignment': 'center-center',
        }

    # Important: first remove the original style/direct border presentation.
    # Some Word table styles otherwise re-introduce grid lines after save/open.
    clear_original_table_style(table)

    # Explicitly disable every table-level border. Cell-level rules below define
    # the only visible lines and survive pagination more reliably.
    tblPr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge_name in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{edge_name}')
        el.set(qn('w:val'), 'none')
        el.set(qn('w:sz'), '0')
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), spec.color)
        borders.append(el)
    tblPr.append(borders)

    # Define exactly three horizontal rules.
    for cell in table.rows[0].cells:
        _set_cell_border(cell, 'top', spec.top_sz, spec.color)
        _set_cell_border(cell, 'bottom', spec.header_sz, spec.color)
    for cell in table.rows[-1].cells:
        _set_cell_border(cell, 'bottom', spec.bottom_sz, spec.color)

    # Default academic layout: all cell content is centered in both axes.
    _center_cell_content(table)

    return {
        'rows': len(table.rows),
        'cols': len(table.columns),
        'spec': spec.to_dict(),
        'original_style_cleared': True,
        'content_alignment': 'center-center',
    }


def convert_tables_to_three_line(doc, baseline_path: str | None = None, include_unknown: bool = False):
    """Format data tables as centered three-line tables; skip layout tables."""
    results = []
    for i, table in enumerate(doc.tables):
        preceding = _preceding_paragraph_text(table)
        cls = classify_table(table, preceding)
        item = {'index': i, 'classification': cls, 'action': 'skipped'}
        if cls['kind'] == 'data' or (include_unknown and cls['kind'] == 'unknown'):
            spec = infer_spec_from_baseline(baseline_path, i) if baseline_path else None
            formatted = make_three_line_table(table, spec or ThreeLineSpec())
            item.update(formatted)
            item['action'] = 'formatted'
        results.append(item)
    return results


# Backward-compatible alias. Layout/unknown tables remain skipped by default.
def convert_all_tables_to_three_line(doc):
    return convert_tables_to_three_line(doc)
