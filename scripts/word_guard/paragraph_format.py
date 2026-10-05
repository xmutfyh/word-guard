# -*- coding: utf-8 -*-
"""Default body-paragraph alignment for Word Guard.

Apply justified alignment to likely running body-text paragraphs while
preserving headings, captions, equations, and explicit center/right alignment.

Many institutional Word templates place their real narrative body inside a
large layout table. This formatter therefore handles top-level paragraphs plus
paragraphs inside tables that are classified as layout tables, while leaving
data/unknown table content to the table formatter.
"""
from __future__ import annotations

import re

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from word_guard.table_format import classify_table, _preceding_paragraph_text


_MATH_TAGS = {qn('m:oMath'), qn('m:oMathPara')}
_SKIP_STYLE_TOKENS = (
    'heading', 'title', 'subtitle', 'caption', 'toc', 'contents',
    'header', 'footer', 'figure', 'table title',
    '标题', '题注', '目录', '页眉', '页脚',
)

_CAPTION_RE = re.compile(
    r'^\s*(?:图|表|figure|fig\.?|table)\s*\d+(?:[.\-–—]\d+|[a-z])?\b',
    re.IGNORECASE,
)
_SECTION_HEADING_RE = re.compile(
    r'^\s*(?:'
    r'[一二三四五六七八九十百]+[、.]'
    r'|\d+(?:\.\d+)*[、.]\s*[^\d]'
    r')'
)
_SENTENCE_PUNCT_RE = re.compile(r'[。！？；.!?;]')


def _contains_math(paragraph) -> bool:
    element = paragraph._p
    if element.tag in _MATH_TAGS:
        return True
    return (
        element.find('.//' + qn('m:oMath')) is not None
        or element.find('.//' + qn('m:oMathPara')) is not None
    )


def _has_outline_level(paragraph) -> bool:
    pPr = paragraph._p.pPr
    return pPr is not None and pPr.find(qn('w:outlineLvl')) is not None


def _style_name(paragraph) -> str:
    try:
        style = paragraph.style
        if style is None:
            return ''
        return ' '.join(filter(None, [getattr(style, 'name', ''), getattr(style, 'style_id', '')])).lower()
    except Exception:
        return ''


def classify_body_paragraph(paragraph) -> tuple[bool, str]:
    """Return ``(is_body, reason)`` for one paragraph."""
    text = paragraph.text.strip()
    if not text:
        return False, 'empty'

    if _contains_math(paragraph):
        return False, 'equation'

    if paragraph.alignment in (WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT):
        return False, 'explicit_center_or_right'

    style_name = _style_name(paragraph)
    if any(token in style_name for token in _SKIP_STYLE_TOKENS):
        return False, 'heading_or_caption_style'

    if _has_outline_level(paragraph):
        return False, 'outline_heading'

    if _CAPTION_RE.match(text):
        return False, 'caption_text'

    if _SECTION_HEADING_RE.match(text) and len(text) <= 80:
        return False, 'section_heading_text'

    # Parenthetical template instructions such as `（...300字以内）` are not
    # running body prose. Numbered items like `（3）...` are unaffected because
    # they do not end with the closing parenthesis.
    if len(text) <= 60 and ((text.startswith('（') and text.endswith('）')) or (text.startswith('(') and text.endswith(')'))):
        return False, 'parenthetical_note'

    # Form prompts/instruction labels commonly end in a colon. Leaving them
    # untouched avoids stretching template labels while normal prose stays justified.
    if len(text) <= 100 and text.endswith(('：', ':')):
        return False, 'form_prompt_or_label'

    # Running body prose is normally long enough to wrap or contains sentence
    # punctuation. Keep short form fields/labels/signature lines untouched.
    compact_len = len(re.sub(r'\s+', '', text))
    # Short numeric/date/form values such as `2002.3` are not body prose.
    # Require enough running text before applying justification.
    if compact_len < 18:
        return False, 'short_label_or_field'

    return True, 'body_text'


def _format_one(paragraph, location: dict) -> dict:
    is_body, reason = classify_body_paragraph(paragraph)
    item = {
        **location,
        'action': 'skipped',
        'reason': reason,
        'text_preview': paragraph.text.strip()[:100],
    }
    if is_body:
        previous = paragraph.alignment
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        item.update({
            'action': 'justified',
            'previous_alignment': int(previous) if previous is not None else None,
            'new_alignment': 'justify',
        })
    return item


def _iter_unique_cell_paragraphs(table):
    """Yield each paragraph in each unique cell once, including nested tables."""
    seen_cells = set()
    for row_idx, row in enumerate(table.rows):
        for col_idx, cell in enumerate(row.cells):
            tc = cell._tc
            if tc in seen_cells:
                continue
            seen_cells.add(tc)
            for p_idx, paragraph in enumerate(cell.paragraphs):
                yield row_idx, col_idx, p_idx, paragraph
            for nested_idx, nested in enumerate(cell.tables):
                for nr, nc, np, paragraph in _iter_unique_cell_paragraphs(nested):
                    yield f'{row_idx}.{nested_idx}.{nr}', nc, np, paragraph


def _is_reference_heading(text: str) -> bool:
    normalized = re.sub(r"\s+", "", text).lower().rstrip('：:')
    return normalized in {'参考文献', 'references', 'bibliography'}


def _format_sequence(paragraphs_with_locations, results):
    """Format one continuous body-text sequence, excluding reference sections."""
    in_references = False
    for paragraph, location in paragraphs_with_locations:
        text = paragraph.text.strip()
        if _is_reference_heading(text):
            in_references = True
            item = _format_one(paragraph, location)
            item['action'] = 'skipped'
            item['reason'] = 'reference_heading'
            results.append(item)
            continue
        if in_references:
            results.append({
                **location,
                'action': 'skipped',
                'reason': 'reference_section',
                'text_preview': text[:100],
            })
            continue
        results.append(_format_one(paragraph, location))


def justify_body_paragraphs(doc) -> list[dict]:
    """Justify likely body text in the document and in layout-table body areas.

    Data/unknown table cells are left untouched here so the table formatter can
    keep its own cell-alignment policy. Layout tables are included because many
    application/report templates use them as page containers for normal prose.
    Reference sections are excluded from this default because they are not body
    prose and often follow their own journal/template alignment rules.
    """
    results = []

    _format_sequence(
        ((paragraph, {'location': 'document', 'paragraph_index': idx})
         for idx, paragraph in enumerate(doc.paragraphs)),
        results,
    )

    for table_idx, table in enumerate(doc.tables):
        cls = classify_table(table, _preceding_paragraph_text(table))
        if cls.get('kind') != 'layout':
            continue

        # Treat every unique cell as its own text sequence. This lets a
        # reference list contained in one template cell be excluded without
        # suppressing later narrative cells in the same layout table.
        seen_cells = set()
        for row_idx, row in enumerate(table.rows):
            for col_idx, cell in enumerate(row.cells):
                tc = cell._tc
                if tc in seen_cells:
                    continue
                seen_cells.add(tc)
                seq = (
                    (paragraph, {
                        'location': 'layout_table',
                        'table_index': table_idx,
                        'row': row_idx,
                        'column': col_idx,
                        'cell_paragraph_index': p_idx,
                    })
                    for p_idx, paragraph in enumerate(cell.paragraphs)
                )
                _format_sequence(seq, results)

                # Nested layout-table paragraphs are uncommon; handle them
                # independently so they do not inherit reference-state flags.
                for nested_idx, nested in enumerate(cell.tables):
                    nested_cls = classify_table(nested, '')
                    if nested_cls.get('kind') != 'layout':
                        continue
                    for nr, nc, np, paragraph in _iter_unique_cell_paragraphs(nested):
                        results.append(_format_one(paragraph, {
                            'location': 'layout_table_nested',
                            'table_index': table_idx,
                            'nested_table_index': nested_idx,
                            'row': nr,
                            'column': nc,
                            'cell_paragraph_index': np,
                        }))

    return results

